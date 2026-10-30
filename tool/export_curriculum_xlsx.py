"""Build ONE Excel file that shows the whole curriculum, level by level.

    python tool/export_curriculum_xlsx.py <dump.json> <out.xlsx>

<dump.json> comes from the app's own lesson data (a throwaway test prints it),
so the workbook says what the app really contains, not what a plan hoped for.
"""

import json
import sys
from collections import Counter, OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src, out = sys.argv[1], sys.argv[2]
data = json.load(open(src, encoding='utf-8'))
rows = data['rows']
letters = data['letters']

LEVELS = {1: 'Ages 3-4', 2: 'Ages 5-6', 3: 'Ages 7-9', 4: 'Ages 10-12'}
WORLD_ORDER = ['english', 'telugu', 'hindi', 'tamil', 'numbers']
WORLD_NAME = {'english': 'English Forest', 'telugu': 'Telugu Kingdom', 'hindi': 'Hindi Valley',
              'tamil': 'Tamil Garden', 'numbers': 'Number Galaxy'}
ALPHA_TITLE = {'english': 'ABC (A to Z)', 'telugu': 'Telugu letters', 'hindi': 'Hindi letters',
               'tamil': 'Tamil letters'}
LANG = {'en': 'English', 'te': 'Telugu', 'hi': 'Hindi', 'ta': 'Tamil'}
KIND = {'alphabet': 'Letter / word quiz', 'words': 'Word quiz', 'numbers': 'Numbers', 'story': 'Story + question',
        'tracing': 'Writing (trace the letter)', 'game': 'Game'}
STEP = {'listen': 'listen', 'pickLetter': 'pick the letter', 'pickWord': 'pick the picture word',
        'pickQuantity': 'pick the quantity', 'count': 'count', 'compare': 'compare', 'choice': 'quiz question',
        'write': 'write'}

head_fill = PatternFill('solid', fgColor='4B3DBA')
head_font = Font(bold=True, color='FFFFFF')
warn_fill = PatternFill('solid', fgColor='FFE9C2')
ok_fill = PatternFill('solid', fgColor='DFF3E4')
lvl_fill = {1: 'E3F6E9', 2: 'E6F0FF', 3: 'F3E8FF', 4: 'FFF1D6'}
thin = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, widths):
    for c, w in enumerate(widths, 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = 'A2'


def finish(ws):
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = border
    ws.auto_filter.ref = ws.dimensions


base = {(r['level'], r['id']) for r in rows if not r['preview']}
full = [r for r in rows if r['preview']]


def language(r):
    code = r['locale'].split('-')[0]
    return LANG.get(code, code)


def what(r):
    if r['steps']:
        return ' > '.join(STEP.get(s, s) for s in r['steps'])
    return KIND.get(r['kind'], r['kind'])


def writes(r):
    return r['kind'] == 'tracing' or 'write' in r['steps']


wb = Workbook()

# ---- Overview ---------------------------------------------------------------
ws = wb.active
ws.title = 'Overview'
ws.append(['Level', 'Ages'] + [WORLD_NAME[w] for w in WORLD_ORDER] + ['Total'])
counts = Counter((r['level'], r['worldId']) for r in full)
base_counts = Counter((r['level'], r['worldId']) for r in rows if not r['preview'])
for lv in LEVELS:
    vals = [counts[(lv, w)] for w in WORLD_ORDER]
    ws.append([f'Level {lv}', LEVELS[lv]] + vals + [sum(vals)])
ws.append([])
ws.append(['Quests the child sees WITHOUT the parent switch "Try the new lessons"'])
ws.append(['Level', 'Ages'] + [WORLD_NAME[w] for w in WORLD_ORDER] + ['Total'])
for lv in LEVELS:
    vals = [base_counts[(lv, w)] for w in WORLD_ORDER]
    ws.append([f'Level {lv}', LEVELS[lv]] + vals + [sum(vals)])
ws.append([])
ws.append(['Letters taught in "All letters" (Level 1 writing):'] + [''] * 1 +
          [f"{len(letters[w])}" if w in letters else '-' for w in WORLD_ORDER])
ws.append([])
notes = [
    'HOW TO READ THIS FILE',
    'Each Level sheet lists every quest in that level, world by world, in the order a child meets them.',
    '"Needs preview switch" = Yes means the quest is only shown after a parent turns on "Try the new lessons".',
    '"Writing" = Yes means the child traces or writes. Writing is only in Level 1 (all letters of every language plus numbers 0-10).',
    'Level 1 also has an "All letters" card per language (see the Level 1 sheet and the Level 1 letters sheet).',
    '',
    'WHAT IS STILL THIN (honest status)',
    'Levels 3 and 4 have only one story quest per language world and two in Number Galaxy: the curriculum pack holds the material only as outlines.',
    'Telugu, Hindi and Tamil letters have no stroke order (only a letter outline); English capitals and numbers 0-10 have a stroke route that no teacher has verified.',
    'No recorded audio exists; the phone voice reads every line. Telugu, Hindi and Tamil sentences written for the app are drafts for a native reader.',
    'Syllables / letter combinations (matras) and spelling are not playable yet: the pack records are blueprints.',
]
for n in notes:
    ws.append([n])
for c in range(1, 9):
    ws.column_dimensions[get_column_letter(c)].width = 18 if c > 2 else 12
ws.column_dimensions['A'].width = 16
for r in (1, 8):
    for cell in ws[r]:
        cell.fill = head_fill
        cell.font = head_font
for r in range(1, ws.max_row + 1):
    if ws.cell(row=r, column=1).value in ('HOW TO READ THIS FILE', 'WHAT IS STILL THIN (honest status)'):
        ws.cell(row=r, column=1).font = Font(bold=True)
    if ws.cell(row=r, column=1).value == 'Quests the child sees WITHOUT the parent switch "Try the new lessons"':
        ws.cell(row=r, column=1).font = Font(bold=True)

# ---- one sheet per level ----------------------------------------------------
HEAD = ['World', '#', 'Lesson id', 'Title', 'What the child does', 'Objective', 'Topic / skill',
        'Examples', 'Needs first', 'Language', 'Writing', 'Needs preview switch', 'Pictures still missing']
WIDTHS = [16, 5, 22, 28, 40, 44, 26, 26, 22, 10, 9, 12, 22]
for lv in LEVELS:
    ws = wb.create_sheet(f'Level {lv} ({LEVELS[lv]})')
    ws.append(HEAD)
    style_header(ws, WIDTHS)
    fill = PatternFill('solid', fgColor=lvl_fill[lv])
    # an "All letters" card leads the trail for Level 1 languages
    for w in WORLD_ORDER:
        if lv == 1 and w in letters:
            ws.append([WORLD_NAME[w], 0, f'alpha-{w}-001..{len(letters[w]):03d}', 'All letters',
                       f'See, hear and trace each of the {len(letters[w])} letters, then Next',
                       'Learn every letter from scratch', ALPHA_TITLE[w], '', '', LANG.get(
                           {'english': 'en', 'telugu': 'te', 'hindi': 'hi', 'tamil': 'ta'}[w], ''),
                       'Yes', 'No', ''])
        n = 0
        for r in [x for x in full if x['level'] == lv and x['worldId'] == w]:
            n += 1
            ws.append([
                WORLD_NAME[w], n, r['id'], r['title'], what(r), r['objective'],
                ' / '.join(x for x in [r['topic'], r['skill']] if x),
                ', '.join(r['examples']), ', '.join(r['prerequisites']), language(r),
                'Yes' if writes(r) else 'No',
                'No' if (lv, r['id']) in base else 'Yes',
                ', '.join(r['missingImages']),
            ])
    for row in ws.iter_rows(min_row=2):
        row[0].fill = fill
        if row[10].value == 'Yes':
            row[10].fill = ok_fill
        if row[11].value == 'Yes':
            row[11].fill = warn_fill
    finish(ws)

# ---- letters of Level 1 ------------------------------------------------------
ws = wb.create_sheet('Level 1 letters')
ws.append(['Language', '#', 'Letter', 'Picture word', 'Meaning', 'Picture', 'Stroke route (order & direction)'])
style_header(ws, [12, 5, 10, 16, 16, 40, 34])
for w in ['english', 'telugu', 'hindi', 'tamil']:
    for i, l in enumerate(letters[w], 1):
        route = 'Yes, unverified' if w == 'english' else 'No (letter outline only)'
        ws.append([WORLD_NAME[w].split()[0], i, l['letter'], l['word'] or '(letter only)', l['meaning'],
                   l['picture'].replace('assets/images/', '') if l['picture'] else
                   ('MISSING' if l['word'] else ''), route])
for row in ws.iter_rows(min_row=2):
    if row[5].value == 'MISSING':
        row[5].fill = warn_fill
finish(ws)

wb.save(out)
print('saved', out, {s.title: s.max_row - 1 for s in wb})
